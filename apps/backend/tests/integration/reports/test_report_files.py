import asyncio
import io
import uuid

import pytest
from docx import Document
from pypdf import PdfReader
from psycopg.types.json import Jsonb

from email_agentic.agents.contracts import Draft
from email_agentic.bootstrap import build_container
from email_agentic.domains.reports.service import ReportEdit
from email_agentic.domains.research.service import ResearchRequest
from email_agentic.integrations.report_export import render_export
from email_agentic.platform.config import Settings
from email_agentic.platform.errors import Conflict, NotFound

MARKDOWN = "# Hồ sơ doanh nghiệp\n\nNghiên cứu và đối tác. **Thông tin** [Nguồn](https://example.com/source).\n\n| Mục | Nội dung |\n|---|---|\n| Sản phẩm | Phân tích tài liệu |\n"


async def test_real_document_exports_preserve_vietnamese_tables_and_sources():
    pdf = await render_export(MARKDOWN, "pdf", "Báo cáo")
    reader = PdfReader(io.BytesIO(pdf))
    assert "Nghiên cứu" in "\n".join(page.extract_text() for page in reader.pages)
    assert any(
        a.get_object().get("/A", {}).get("/URI") == "https://example.com/source"
        for page in reader.pages
        for a in page.get("/Annots", [])
    )
    docx = Document(io.BytesIO(await render_export(MARKDOWN, "docx", "Báo cáo")))
    assert any("Nghiên cứu" in p.text for p in docx.paragraphs)
    assert docx.tables[0].cell(1, 1).text == "Phân tích tài liệu"
    assert any(r.target_ref == "https://example.com/source" for r in docx.part.rels.values())
    assert await render_export(MARKDOWN, "md", "Báo cáo") == MARKDOWN.encode()


async def test_untrusted_html_images_and_local_resources_cannot_be_loaded(tmp_path):
    secret = tmp_path / "private.txt"
    secret.write_text("PRIVATE_LOCAL_CONTENT")
    data = await render_export(
        f'# Safe\n\n<img src="{secret.as_uri()}">\n\n![image]({secret.as_uri()})\n\n<script>run()</script>',
        "pdf",
        "Safe",
    )
    assert "PRIVATE_LOCAL_CONTENT" not in "".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)


async def make_report(c, owner):
    rid = await c.research.create(
        owner, ResearchRequest(company_name="Example", domain="example.com"), str(uuid.uuid4())
    )
    run = await c.research.get(owner, rid)
    await c.reports.initial(
        run, Draft(markdown=MARKDOWN, email_subject="Trao đổi", email_body="Cảm ơn anh/chị."), {}
    )
    return rid, await c.reports.latest(owner, rid)


async def test_selected_files_are_versioned_owned_and_checked_before_approval(db):
    owner = f"{uuid.uuid4()}@example.invalid"
    async with build_container(Settings(APP_ENV="test", DATABASE_URL=db.pool.conninfo, _env_file=None)) as c:
        rid, initial = await make_report(c, owner)
        assert initial["artifacts"] == []
        exported = await c.reports.files.export(owner, rid, initial["id"], "pdf")
        assert (await c.reports.files.export(owner, rid, initial["id"], "pdf")) == exported
        with pytest.raises(NotFound):
            await c.reports.files.get("other@example.invalid", uuid.UUID(exported["id"]))
        edit = ReportEdit(
            version=1,
            markdown=MARKDOWN,
            email_subject="Trao đổi",
            email_body="Cảm ơn anh/chị.",
            attachment_formats=["pdf", "docx"],
        )
        await c.reports.edit(owner, rid, edit)
        selected = await c.reports.latest(owner, rid)
        assert selected["version"] == 2 and len(selected["artifacts"]) == 2
        assert initial["artifacts"] == []
        approval = await db.one("SELECT id FROM approvals WHERE report_id=%s", (selected["id"],))
        await db.execute(
            "UPDATE report_files SET data=%s WHERE id=%s",
            (b"changed", uuid.UUID(selected["artifacts"][0]["id"])),
        )
        with pytest.raises(Conflict, match="attachment"):
            await c.approvals.decide(owner, approval["id"], True)
        assert (await db.one("SELECT status FROM approvals WHERE id=%s", (approval["id"],)))[
            "status"
        ] == "pending"


async def test_unchanged_attachment_is_reused_when_enabling_send_and_consumes_matching_draft(db):
    from email_agentic.domains.reports.service import Recipients

    async with build_container(Settings(APP_ENV="test", DATABASE_URL=db.pool.conninfo, _env_file=None)) as c:
        owner = f"{uuid.uuid4()}@example.invalid"
        rid, initial = await make_report(c, owner)
        await db.execute(
            "UPDATE research_runs SET source_message=%s WHERE id=%s",
            (
                Jsonb(
                    {"thread_id": "thread-1", "headers": {"Message-ID": "<source@example.invalid>"}}
                ),
                rid,
            ),
        )
        edit = dict(
            markdown=MARKDOWN,
            email_subject=initial["email_subject"],
            email_body=initial["email_body"],
            recipients=Recipients(to=["partner@example.invalid"]),
            attachment_formats=["pdf"],
        )
        await c.reports.edit(owner, rid, ReportEdit(version=1, **edit))
        selected = await c.reports.latest(owner, rid)
        draft_id = await c.drafts.request(owner, rid, 2)
        draft_package = await c.drafts.package(owner, draft_id)
        assert draft_package["attachments"][0]["data"].startswith(b"%PDF")
        await c.drafts.complete(draft_id, "succeeded", "provider-draft", "thread-1")
        await c.reports.edit(owner, rid, ReportEdit(version=2, send_email=True, **edit))
        latest = await c.reports.latest(owner, rid)
        assert latest["artifacts"] == selected["artifacts"]
        approval = await db.one("SELECT id FROM approvals WHERE report_id=%s", (latest["id"],))
        await c.approvals.decide(owner, approval["id"], True)
        action = await db.one(
            "SELECT id FROM publication_actions WHERE approval_id=%s AND kind='email'", (approval["id"],)
        )
        package = await c.publication.package(owner, action["id"])
        assert package["gmail_draft_provider_id"] == "provider-draft"
        assert package["attachments"] == draft_package["attachments"]
        await c.publication.complete(action["id"], "succeeded", "sent-1")
        assert (await db.one("SELECT status FROM gmail_drafts WHERE id=%s", (draft_id,)))["status"] == "sent"


@pytest.mark.parametrize("change", ["edit", "raw_expiry", "history_expiry"])
async def test_edit_renders_without_run_lock_and_revalidates_before_saving(db, monkeypatch, change):
    import email_agentic.domains.reports.files as files

    started, release = asyncio.Event(), asyncio.Event()

    async def controlled_render(markdown, format, title):
        started.set()
        await release.wait()
        return markdown.encode()

    monkeypatch.setattr(files, "render_export", controlled_render)
    async with build_container(Settings(APP_ENV="test", DATABASE_URL=db.pool.conninfo, _env_file=None)) as c:
        owner = f"{uuid.uuid4()}@example.invalid"
        rid, report = await make_report(c, owner)
        edit = ReportEdit(version=1, markdown="New report", email_subject=report["email_subject"],
                          email_body=report["email_body"], attachment_formats=["md"])
        task = asyncio.create_task(c.reports.edit(owner, rid, edit))
        try:
            await asyncio.wait_for(started.wait(), 5)
            async with db.transaction() as connection:
                await connection.execute("SELECT id FROM research_runs WHERE id=%s FOR UPDATE NOWAIT", (rid,))
            if change == "edit":
                await c.reports.edit(owner, rid, edit.model_copy(update={"attachment_formats": []}))
            else:
                column = "raw_purged_at" if change == "raw_expiry" else "history_purged_at"
                await db.execute(f"UPDATE research_runs SET {column}=now() WHERE id=%s", (rid,))
            release.set()
            with pytest.raises(NotFound if change == "history_expiry" else Conflict):
                await asyncio.wait_for(task, 5)
            assert not await db.all("SELECT f.id FROM report_files f JOIN reports r ON r.id=f.report_id WHERE r.run_id=%s", (rid,))
            assert (await c.reports.latest(owner, rid))["version"] == (2 if change == "edit" else 1)
        finally:
            if not task.done():
                task.cancel()
            await asyncio.gather(task, return_exceptions=True)


async def test_export_does_not_recreate_content_purged_during_render(db, monkeypatch):
    import email_agentic.domains.reports.files as files

    started, release = asyncio.Event(), asyncio.Event()

    async def controlled_render(markdown, format, title):
        started.set()
        await release.wait()
        return markdown.encode()

    monkeypatch.setattr(files, "render_export", controlled_render)
    async with build_container(Settings(APP_ENV="test", DATABASE_URL=db.pool.conninfo, _env_file=None)) as c:
        owner = f"{uuid.uuid4()}@example.invalid"
        rid, report = await make_report(c, owner)
        task = asyncio.create_task(c.reports.files.export(owner, rid, report["id"], "md"))
        try:
            await asyncio.wait_for(started.wait(), 5)
            async with db.transaction() as connection:
                await connection.execute("SELECT id FROM research_runs WHERE id=%s FOR UPDATE NOWAIT", (rid,))
                await connection.execute("UPDATE research_runs SET history_purged_at=now() WHERE id=%s", (rid,))
            release.set()
            with pytest.raises(NotFound):
                await asyncio.wait_for(task, 5)
            assert not await db.all("SELECT id FROM report_files WHERE report_id=%s", (report["id"],))
        finally:
            if not task.done():
                task.cancel()
            await asyncio.gather(task, return_exceptions=True)
