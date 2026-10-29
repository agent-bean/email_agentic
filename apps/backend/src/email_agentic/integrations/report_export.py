"""Bounded document rendering with no network or arbitrary filesystem resources."""

import asyncio
import io
import json
import sys
from html import escape

MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "md": "text/markdown; charset=utf-8",
}
MAX_BYTES = 10 * 1024 * 1024
RENDER_SLOTS = asyncio.Semaphore(2)


def render(markdown, format, title):
    from markdown_it import MarkdownIt

    parser = MarkdownIt("commonmark", {"html": False}).enable("table")
    if format == "md":
        return markdown.encode()
    if format == "pdf":
        from weasyprint import HTML
        from weasyprint.urls import URLFetcher

        class NoResources(URLFetcher):
            def fetch(self, url, headers=None):
                raise ValueError("External resources are disabled")

        # Images are excluded, including data URIs; links remain clickable but are never fetched.
        parser.disable("image")
        css = """@page { size:A4; margin:20mm; @bottom-right {content:counter(page)} }
        body {font-family:"DejaVu Sans",sans-serif;font-size:10pt;line-height:1.5;color:#17202a}
        h1,h2,h3 {break-after:avoid} h1{font-size:20pt} h2{font-size:15pt} h3{font-size:12pt}
        table{border-collapse:collapse;width:100%;font-size:9pt} td,th{border:1px solid #ccd3da;padding:5px;overflow-wrap:anywhere}
        pre{white-space:pre-wrap} a{color:#155e75;overflow-wrap:anywhere}"""
        return HTML(
            string='<html lang="vi"><head><meta charset="utf-8"><title>'
            + escape(title)
            + "</title><style>"
            + css
            + "</style></head><body>"
            + parser.render(markdown)
            + "</body></html>",
            url_fetcher=NoResources(),
        ).write_pdf()
    if format != "docx":
        raise ValueError("Unsupported format")
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    document = Document()
    document.core_properties.title = title
    document.styles["Normal"].font.name = "DejaVu Sans"
    document.styles["Normal"].font.size = Pt(10)

    def inline(paragraph, children):
        bold, italic, href = False, False, None
        for token in children or []:
            if token.type == "strong_open":
                bold = True
            elif token.type == "strong_close":
                bold = False
            elif token.type == "em_open":
                italic = True
            elif token.type == "em_close":
                italic = False
            elif token.type == "link_open":
                href = token.attrGet("href")
            elif token.type == "link_close":
                href = None
            elif token.type in {"softbreak", "hardbreak"}:
                paragraph.add_run().add_break()
            elif token.type in {"text", "code_inline", "image"}:
                if href and href.startswith(("https://", "http://")):
                    link = OxmlElement("w:hyperlink")
                    link.set(
                        qn("r:id"),
                        paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True),
                    )
                    run = OxmlElement("w:r")
                    text = OxmlElement("w:t")
                    text.text = token.content
                    run.append(text)
                    link.append(run)
                    paragraph._p.append(link)
                else:
                    run = paragraph.add_run(token.content)
                    run.bold = bold
                    run.italic = italic

    tokens = parser.parse(markdown)
    paragraph = None
    table = None
    cells = None
    column = 0
    list_depth = 0
    for index, token in enumerate(tokens):
        if token.type == "heading_open":
            paragraph = document.add_heading(level=min(6, int(token.tag[1])))
        elif token.type in {"bullet_list_open", "ordered_list_open"}:
            list_depth += 1
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            list_depth -= 1
        elif token.type == "paragraph_open":
            paragraph = document.add_paragraph(style="List Bullet" if list_depth else None)
        elif token.type == "table_open":
            table = document.add_table(rows=0, cols=0)
            table.style = "Table Grid"
        elif token.type == "tr_open":
            cells = None
            column = 0
        elif token.type in {"th_open", "td_open"}:
            if cells is None:
                # Header width is obtained from its row; Word tables require columns first.
                if not table.columns:
                    start = index
                    width = 0
                    for item in tokens[start:]:
                        if item.type == "tr_close":
                            break
                        if item.type in {"th_open", "td_open"}:
                            width += 1
                    for _ in range(width):
                        table.add_column(Pt(450 / max(1, width)))
                cells = table.add_row().cells
            paragraph = cells[min(column, len(cells) - 1)].paragraphs[0]
            column += 1
        elif token.type == "inline":
            if paragraph is None:
                paragraph = document.add_paragraph()
            inline(paragraph, token.children)
        elif token.type in {"fence", "code_block"}:
            document.add_paragraph(token.content, style="No Spacing")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


async def render_export(markdown, format, title):
    if format not in MIME or len(markdown) > 100000:
        raise ValueError("Invalid document input")
    async with RENDER_SLOTS:
        return await _render_bounded(markdown, format, title)


async def _render_bounded(markdown, format, title):
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        "-m",
        "email_agentic.integrations.report_export",
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
    )
    try:
        async with asyncio.timeout(30):
            data, _ = await process.communicate(json.dumps([markdown, format, title]).encode())
        if process.returncode or not data or len(data) > MAX_BYTES:
            raise ValueError("Document rendering failed or exceeded limits")
        return data
    finally:
        if process.returncode is None:
            process.kill()
        await process.wait()


if __name__ == "__main__":
    import resource

    resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (25, 25))
    sys.stdout.buffer.write(render(*json.loads(sys.stdin.buffer.read(700000))))
